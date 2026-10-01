"""Spreadsheet resources for /resources, generated with openpyxl.

Called from build.py. Produces:
  marketing-budget-template.xlsx
  media-list-template.xlsx
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

INK = "111111"
PAPER = "F7F5EF"
LIME = "D4F25A"
LIME_SOFT = "EEF9C4"
RULE = "D9D5C8"
INPUT = "FFFDE7"

F_TITLE = Font(name="Georgia", size=20, color=INK)
F_H = Font(name="Calibri", size=11, bold=True, color=PAPER)
F_B = Font(name="Calibri", size=11, bold=True, color=INK)
F_N = Font(name="Calibri", size=11, color=INK)
F_MUTED = Font(name="Calibri", size=10, italic=True, color="6B6A63")
FILL_H = PatternFill("solid", fgColor=INK)
FILL_LIME = PatternFill("solid", fgColor=LIME)
FILL_SOFT = PatternFill("solid", fgColor=LIME_SOFT)
FILL_INPUT = PatternFill("solid", fgColor=INPUT)
THIN = Side(style="thin", color=RULE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
MONEY = '"$"#,##0;[Red]-"$"#,##0;"-"'
MONEY2 = '"$"#,##0.00'
PCT = "0.0%"
FOOTER = "Free resource · shilikajain.com/resources · Book a 30-minute teardown: shilikajain.com/contact"


def _brand(ws, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws["A2"] = subtitle
    ws["A2"].font = F_MUTED
    ws.sheet_view.showGridLines = False


def _header(ws, row: int, headers: list[str], widths: list[int] | None = None) -> None:
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = F_H
        c.fill = FILL_H
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 32


def _instructions(ws, title: str, lines: list[tuple[str, str]]) -> None:
    _brand(ws, title, "Shilika Jain · shilikajain.com · " + "free template")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100
    r = 4
    for head, text in lines:
        a = ws.cell(row=r, column=1, value=head)
        b = ws.cell(row=r, column=2, value=text)
        a.font = F_B
        b.font = F_N
        a.alignment = WRAP
        b.alignment = WRAP
        if head:
            a.fill = FILL_SOFT
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Want help?").font = F_B
    ws.cell(row=r, column=1).fill = FILL_LIME
    ws.cell(row=r, column=2, value=(
        "Want a senior operator on your launch? Book a 30-minute teardown: shilikajain.com/contact. "
        "Interactive planners and calculators: shilikajain.com/tools"
    )).font = F_N
    ws.cell(row=r, column=2).alignment = WRAP


# --------------------------------------------------------------------------
# Marketing budget template
# --------------------------------------------------------------------------

CHANNELS = [
    "SEO (content, technical fixes)",
    "GEO / AI search (citable content, listings)",
    "Google Ads",
    "Meta Ads",
    "LinkedIn Ads",
    "Reddit (ads, community time)",
    "X / Twitter (ads, content)",
    "LinkedIn organic (content support)",
    "YouTube (production, ads)",
    "TikTok / Shorts / Reels",
    "Email (platform sends, content)",
    "WhatsApp (Business API, broadcasts)",
    "Communities (Discord, Slack, Telegram)",
    "Product Hunt and launch sites",
    "PR / earned media (content, data studies)",
    "Podcasts (guesting, production)",
    "Influencers / KOLs",
    "Partnerships and integrations",
    "Events and conferences (as a channel test)",
    "Referral program (incentives)",
]

# (section, [(line item, low, mid, high, note)])
BUDGET = [
    ("People", [
        ("Marketing hire(s), fully loaded", 0, 9000, 16000, "Salary plus benefits and taxes, monthly"),
        ("Fractional marketing or growth lead", 0, 6000, 10000, "Part-time senior operator"),
        ("Freelancers: writing, design, video", 1000, 3000, 6000, ""),
        ("Founder time (track hours, not cash)", 0, 0, 0, "Optional: put an opportunity cost here"),
    ]),
    ("Programs by channel", [(c, 0, 0, 0, "") for c in CHANNELS]),
    ("Tools", [
        ("CRM", 0, 100, 500, ""),
        ("Email platform", 0, 100, 400, ""),
        ("Analytics and attribution", 0, 100, 400, ""),
        ("SEO and AI-visibility tools", 0, 150, 500, ""),
        ("Design and video tools", 0, 50, 150, ""),
        ("Social scheduling and listening", 0, 50, 300, ""),
        ("Media database or monitoring", 0, 0, 600, ""),
    ]),
    ("Events", [
        ("Conference tickets and travel", 0, 1500, 5000, "Spread across months or enter in the event month"),
        ("Sponsorships and booths", 0, 0, 10000, ""),
        ("Side events, dinners, meetups", 0, 1000, 4000, ""),
    ]),
    ("PR", [
        ("Fractional PR retainer", 0, 5000, 12000, "Typical fractional range $5K to $12K per month"),
        ("Launch sprint (one-off, enter in launch months)", 0, 0, 0, "Typical range $15K to $40K per sprint"),
        ("Wire distribution", 0, 0, 800, "Only when a wire actually adds value"),
        ("Media training, photography, press kit design", 0, 300, 1000, ""),
    ]),
]

# Example values for the channel programs (seed-stage B2B; replace with your own)
CHANNEL_EXAMPLES = {
    "SEO (content, technical fixes)": (500, 1500, 4000),
    "GEO / AI search (citable content, listings)": (0, 500, 2000),
    "Google Ads": (0, 2000, 6000),
    "LinkedIn Ads": (0, 2500, 8000),
    "Email (platform sends, content)": (0, 300, 1000),
    "Communities (Discord, Slack, Telegram)": (0, 300, 1500),
    "PR / earned media (content, data studies)": (0, 500, 2500),
    "Podcasts (guesting, production)": (0, 200, 1000),
}


def build_budget(path: Path) -> Path:
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    _instructions(ins, "Marketing Budget Template", [
        ("What this is", "A 12-month marketing budget with low, mid and high scenarios, a line for each of the 20 channels in the 20 Marketing Channels Playbook, a channel test tracker and a unit economics calculator."),
        ("Yellow cells", "Inputs. Everything else is calculated. Example numbers are illustrative placeholders for a seed-stage B2B company; replace all of them with your own."),
        ("1. Monthly Budget", "Enter a monthly Low, Mid and High amount per line. Pick the scenario in cell C3. Months 1 to 12 pull the chosen scenario by default; overwrite any month cell to model ramps, launches or one-off costs (for example a launch sprint or a conference)."),
        ("2. Contingency", "Set the contingency rate in cell C4 (10% is a sensible default for early-stage teams). It is applied to the total of all sections."),
        ("3. Channel Tests", "Before spending on a new channel, write the hypothesis, test budget and kill criteria. Update spend and outcomes weekly. Cost per outcome calculates automatically. Decide: Kill, Iterate or Scale."),
        ("4. Unit Economics", "Enter spend, new customers, revenue per account, gross margin and churn to calculate CAC, LTV, LTV:CAC and CAC payback. Use these to decide how much each channel is allowed to cost."),
        ("Rule of three", "Run no more than three acquisition channels at once until one works repeatably."),
        ("Note", "Ranges for PR engagements reflect Shilika Jain's own pricing (fractional PR retainer $5K to $12K per month; launch sprint $15K to $40K). All other figures are placeholders."),
    ])

    ws = wb.create_sheet("Monthly Budget")
    _brand(ws, "Monthly Budget", "12-month plan · scenario-driven · yellow cells are inputs")
    ws["B3"] = "Scenario"
    ws["B3"].font = F_B
    ws["C3"] = "Mid"
    ws["C3"].fill = FILL_INPUT
    ws["C3"].font = F_B
    ws["C3"].border = BORDER
    dv = DataValidation(type="list", formula1='"Low,Mid,High"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("C3")
    ws["B4"] = "Contingency rate"
    ws["B4"].font = F_B
    ws["C4"] = 0.10
    ws["C4"].number_format = PCT
    ws["C4"].fill = FILL_INPUT
    ws["C4"].border = BORDER

    headers = ["Section", "Line item", "Low / mo", "Mid / mo", "High / mo"] + [f"Month {i}" for i in range(1, 13)] + ["12-mo total", "Notes"]
    hr = 6
    _header(ws, hr, headers, [18, 44, 11, 11, 11] + [11] * 12 + [14, 40])
    ws.freeze_panes = ws.cell(row=hr + 1, column=3)

    r = hr + 1
    subtotal_rows = []
    for section, items in BUDGET:
        start = r
        for name, lo, mid, hi, note in items:
            if section == "Programs by channel":
                lo, mid, hi = CHANNEL_EXAMPLES.get(name, (0, 0, 0))
            ws.cell(row=r, column=1, value=section).font = F_MUTED
            ws.cell(row=r, column=2, value=name).font = F_N
            for col, v in zip((3, 4, 5), (lo, mid, hi)):
                c = ws.cell(row=r, column=col, value=v)
                c.number_format = MONEY
                c.fill = FILL_INPUT
            for m in range(12):
                col = 6 + m
                c = ws.cell(row=r, column=col, value=f'=IF($C$3="Low",$C{r},IF($C$3="High",$E{r},$D{r}))')
                c.number_format = MONEY
            t = ws.cell(row=r, column=18, value=f"=SUM(F{r}:Q{r})")
            t.number_format = MONEY
            t.font = F_B
            ws.cell(row=r, column=19, value=note).font = F_MUTED
            for col in range(1, 20):
                ws.cell(row=r, column=col).border = BORDER
            r += 1
        end = r - 1
        ws.cell(row=r, column=2, value=f"Subtotal: {section}").font = F_B
        for col in range(3, 19):
            L = get_column_letter(col)
            c = ws.cell(row=r, column=col, value=f"=SUM({L}{start}:{L}{end})")
            c.number_format = MONEY
            c.font = F_B
        for col in range(1, 20):
            ws.cell(row=r, column=col).fill = FILL_SOFT
            ws.cell(row=r, column=col).border = BORDER
        subtotal_rows.append(r)
        r += 2

    # Contingency
    ws.cell(row=r, column=1, value="Contingency").font = F_MUTED
    ws.cell(row=r, column=2, value="Contingency (rate in C4 x all sections)").font = F_B
    for col in range(3, 19):
        L = get_column_letter(col)
        refs = "+".join(f"{L}{sr}" for sr in subtotal_rows)
        c = ws.cell(row=r, column=col, value=f"=$C$4*({refs})")
        c.number_format = MONEY
    for col in range(1, 20):
        ws.cell(row=r, column=col).border = BORDER
    cont_row = r
    r += 2

    ws.cell(row=r, column=2, value="TOTAL MARKETING BUDGET").font = F_B
    for col in range(3, 19):
        L = get_column_letter(col)
        refs = "+".join(f"{L}{sr}" for sr in subtotal_rows + [cont_row])
        c = ws.cell(row=r, column=col, value=f"={refs}")
        c.number_format = MONEY
        c.font = F_B
    for col in range(1, 20):
        ws.cell(row=r, column=col).fill = FILL_LIME
        ws.cell(row=r, column=col).border = BORDER
    r += 2
    ws.cell(row=r, column=2, value=FOOTER).font = F_MUTED

    # Channel tests
    ct = wb.create_sheet("Channel Tests")
    _brand(ct, "Channel Tests", "Write the kill criteria before you spend. Update weekly.")
    heads = ["Channel", "Hypothesis", "Start date", "End date", "Test budget", "Spent to date", "Budget left",
             "Success metric (outcome)", "Target cost per outcome", "Kill criteria", "Outcomes to date",
             "Cost per outcome", "Status vs target", "Decision", "Learnings / notes"]
    _header(ct, 4, heads, [26, 40, 12, 12, 13, 13, 13, 24, 14, 36, 12, 14, 16, 12, 40])
    ct.freeze_panes = "B5"
    examples = [
        ("LinkedIn Ads", "VP Finance at 200 to 1,000 staff SaaS will download a benchmark report", 3000, 1850, "Qualified lead", 150, "Cost per qualified lead above $300 after $2,000 spent", 9),
        ("Google Ads", "Competitor and category terms convert to demo requests from landing page v2", 2500, 2500, "Demo request", 250, "Fewer than 5 demos after full budget", 6),
        ("Podcasts", "Founder guesting on 4 niche shows drives self-reported signups", 0, 0, "Self-reported signup", 0, "Fewer than 3 self-reported mentions after 4 episodes", 0),
    ]
    dec = DataValidation(type="list", formula1='"Running,Kill,Iterate,Scale"', allow_blank=True)
    ct.add_data_validation(dec)
    for i in range(40):
        r = 5 + i
        if i < len(examples):
            ch, hyp, bud, spent, metric, target, kill, outc = examples[i]
            vals = {1: ch + " (example)", 2: hyp, 5: bud, 6: spent, 8: metric, 9: target, 10: kill, 11: outc, 14: "Running"}
            for col, v in vals.items():
                ct.cell(row=r, column=col, value=v)
        ct.cell(row=r, column=7, value=f"=IF(E{r}=\"\",\"\",E{r}-F{r})")
        ct.cell(row=r, column=12, value=f"=IFERROR(F{r}/K{r},\"\")")
        ct.cell(row=r, column=13, value=f"=IF(OR(L{r}=\"\",I{r}=\"\",I{r}=0),\"\",IF(L{r}<=I{r},\"On target\",\"Above target\"))")
        for col in (5, 6, 7, 9, 12):
            ct.cell(row=r, column=col).number_format = MONEY
        for col in (3, 4):
            ct.cell(row=r, column=col).number_format = "yyyy-mm-dd"
        for col in range(1, 16):
            c = ct.cell(row=r, column=col)
            c.border = BORDER
            c.alignment = WRAP
            if col in (1, 2, 3, 4, 5, 6, 8, 9, 10, 11, 14, 15):
                c.fill = FILL_INPUT
        dec.add(f"N{r}")
    ct.conditional_formatting.add("M5:M44", CellIsRule(operator="equal", formula=['"On target"'], fill=FILL_LIME))
    ct.cell(row=47, column=1, value=FOOTER).font = F_MUTED

    # Unit economics
    ue = wb.create_sheet("Unit Economics")
    _brand(ue, "Unit Economics", "CAC, LTV, LTV:CAC and payback. Example inputs are illustrative; replace them.")
    ue.column_dimensions["A"].width = 44
    ue.column_dimensions["B"].width = 16
    ue.column_dimensions["C"].width = 70
    rows = [
        ("INPUTS", None, None, "head"),
        ("Marketing spend in period", 30000, "Programs, tools, people attributable to acquisition", "money"),
        ("Sales spend in period", 20000, "Sales salaries and costs attributable to new customers", "money"),
        ("New customers acquired in period", 10, "Count paying customers, not signups", "int"),
        ("Average revenue per account, monthly", 1500, "ARPA. For annual contracts, divide ACV by 12", "money"),
        ("Gross margin", 0.75, "After hosting, inference and support costs. AI products: check inference costs carefully", "pct"),
        ("Monthly revenue churn", 0.02, "Share of MRR lost per month", "pct"),
        ("OUTPUTS", None, None, "head"),
        ("Blended CAC", "=IFERROR((B5+B6)/B7,0)", "(Marketing + sales spend) / new customers", "money"),
        ("Marketing-only CAC", "=IFERROR(B5/B7,0)", "Useful for channel comparisons", "money"),
        ("Monthly gross profit per account", "=B8*B9", "ARPA x gross margin", "money"),
        ("Expected customer lifetime (months)", "=IFERROR(1/B10,0)", "1 / monthly churn. Cap at 60 for planning if churn is very low", "num"),
        ("LTV (gross-margin based)", "=IFERROR(B8*B9/B10,0)", "ARPA x gross margin / monthly churn", "money"),
        ("LTV : CAC", "=IFERROR(B16/B12,0)", "Common rule of thumb: around 3:1 or better. A heuristic, not a law", "ratio"),
        ("CAC payback (months)", "=IFERROR(B12/B14,0)", "Months of gross profit to recover CAC. Many seed teams aim for under 12 to 18", "num"),
        ("Max affordable CAC at 3:1", "=B16/3", "Use as the ceiling when setting channel test targets", "money"),
    ]
    r = 4
    for label, val, note, kind in rows:
        a = ue.cell(row=r, column=1, value=label)
        if kind == "head":
            for col in range(1, 4):
                ue.cell(row=r, column=col).fill = FILL_H
            a.font = F_H
            r += 1
            continue
        a.font = F_N
        b = ue.cell(row=r, column=2, value=val)
        b.font = F_B
        ue.cell(row=r, column=3, value=note).font = F_MUTED
        if isinstance(val, str) and val.startswith("="):
            b.fill = FILL_SOFT
        else:
            b.fill = FILL_INPUT
        b.number_format = {"money": MONEY, "pct": PCT, "int": "0", "num": "0.0", "ratio": '0.0"x"'}[kind]
        for col in range(1, 4):
            ue.cell(row=r, column=col).border = BORDER
        r += 1
    ue.cell(row=r + 1, column=1, value=FOOTER).font = F_MUTED

    for s in wb.worksheets:
        s.sheet_properties.tabColor = LIME if s.title != "Instructions" else INK
    wb.active = 0
    wb.save(path)
    return path


# --------------------------------------------------------------------------
# Media list template
# --------------------------------------------------------------------------

TIERS = "Tier 1,Tier 2,Tier 3,Trade / vertical,Newsletter,Podcast,Regional"
STATUSES = "Not contacted,Researching,Pitched,Followed up,Responded,Interested,Briefed,Covered,Declined,Do not contact"
PITCH_TYPES = "Exclusive,Embargo,Funding,Data / research,Reactive,Op-ed,Podcast,Expert source,Customer story,Trend,Follow-up,Thank you"
OUTCOMES = "No reply,Declined,Interested,Interview booked,Covered,Op-ed accepted,Later / timing"


def build_media_list(path: Path) -> Path:
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    _instructions(ins, "Media List Template", [
        ("What this is", "A working media list and pitch tracker for founders and comms teams. One row per journalist, one row per pitch."),
        ("Rule 1", "Quality over volume. 30 well-researched contacts beat 500 scraped ones. Every row needs a recent relevant article URL before anyone is pitched."),
        ("Rule 2", "One journalist per outlet per story. Use the Status column to make sure two people at the same outlet are not pitched at once."),
        ("Rule 3", "Only mark 'Embargo-reliable?' as Yes after a journalist has honoured an embargo with you. Unknown is the default."),
        ("Rule 4", "Respect opt-outs. Set Status to 'Do not contact' and keep the row so nobody re-adds them."),
        ("Media List sheet", "Tier and Status use dropdowns. Example rows use placeholders only; delete them before you start. Never copy contact details from sources that prohibit it, and keep the file private."),
        ("Pitch tracker sheet", "Log each pitch with type, subject line and date. Follow-up date calculates as 4 days after sending. Follow up once with something new, then stop."),
        ("Privacy", "This file will hold personal data. Store it securely, share it only with people who need it, and delete contacts you no longer work with."),
        ("Pairs with", "Journalist Pitch Email Templates and Press Kit and Launch PR Template at shilikajain.com/resources."),
    ])

    ml = wb.create_sheet("Media List")
    _brand(ml, "Media List", "One row per journalist. Research before you pitch.")
    heads = ["Outlet", "Tier", "Journalist name", "Beat", "Recent relevant article URL", "Email / contact method",
             "X handle", "Last contacted", "Status", "Notes", "Embargo-reliable?"]
    _header(ml, 4, heads, [24, 16, 24, 24, 42, 30, 18, 14, 16, 44, 16])
    ml.freeze_panes = "B5"
    examples = [
        ("Example Outlet", "Tier 1", "[Reporter name]", "[e.g. AI infrastructure, funding]", "[https://example.com/recent-article]", "[Preferred contact method]", "[@handle]", None, "Researching", "[Why this reporter: their angle on a recent story]", "Unknown"),
        ("Example Trade Publication", "Trade / vertical", "[Reporter name]", "[e.g. DeFi security]", "[https://example.com/recent-article]", "[Preferred contact method]", "[@handle]", None, "Not contacted", "[Note]", "Unknown"),
        ("Example Newsletter", "Newsletter", "[Editor name]", "[e.g. developer tools]", "[https://example.com/issue]", "[Submission form or email]", "[@handle]", None, "Not contacted", "[Note]", "Unknown"),
    ]
    last = 304
    dv_tier = DataValidation(type="list", formula1=f'"{TIERS}"', allow_blank=True)
    dv_status = DataValidation(type="list", formula1=f'"{STATUSES}"', allow_blank=True)
    dv_emb = DataValidation(type="list", formula1='"Yes,No,Unknown"', allow_blank=True)
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True)
    for dv in (dv_tier, dv_status, dv_emb, dv_date):
        ml.add_data_validation(dv)
    dv_tier.add(f"B5:B{last}")
    dv_status.add(f"I5:I{last}")
    dv_emb.add(f"K5:K{last}")
    dv_date.add(f"H5:H{last}")
    for i, row in enumerate(examples):
        for col, v in enumerate(row, 1):
            c = ml.cell(row=5 + i, column=col, value=v)
            c.font = F_MUTED
    for r in range(5, last + 1):
        ml.cell(row=r, column=8).number_format = "yyyy-mm-dd"
        for col in range(1, 12):
            ml.cell(row=r, column=col).border = BORDER
            ml.cell(row=r, column=col).alignment = WRAP
    ml.conditional_formatting.add(f"I5:I{last}", CellIsRule(operator="equal", formula=['"Covered"'], fill=FILL_LIME))
    ml.conditional_formatting.add(f"I5:I{last}", CellIsRule(operator="equal", formula=['"Do not contact"'], fill=PatternFill("solid", fgColor="E6E3DA")))
    ml.auto_filter.ref = f"A4:K{last}"

    pt = wb.create_sheet("Pitch Tracker")
    _brand(pt, "Pitch Tracker", "One row per pitch. Follow up once, with something new.")
    heads = ["Date sent", "Story / angle", "Outlet", "Journalist", "Pitch type", "Subject line",
             "Exclusive / embargo terms", "Follow-up due", "Followed up?", "Outcome", "Coverage URL", "Notes"]
    _header(pt, 4, heads, [13, 34, 22, 22, 16, 40, 22, 14, 13, 18, 36, 36])
    pt.freeze_panes = "C5"
    dv_type = DataValidation(type="list", formula1=f'"{PITCH_TYPES}"', allow_blank=True)
    dv_out = DataValidation(type="list", formula1=f'"{OUTCOMES}"', allow_blank=True)
    dv_yn = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    for dv in (dv_type, dv_out, dv_yn):
        pt.add_data_validation(dv)
    plast = 204
    dv_type.add(f"E5:E{plast}")
    dv_out.add(f"J5:J{plast}")
    dv_yn.add(f"I5:I{plast}")
    pt.cell(row=5, column=2, value="[Example: launch of product X, angle Y]").font = F_MUTED
    pt.cell(row=5, column=3, value="Example Outlet").font = F_MUTED
    pt.cell(row=5, column=4, value="[Reporter name]").font = F_MUTED
    pt.cell(row=5, column=5, value="Embargo").font = F_MUTED
    pt.cell(row=5, column=6, value="[Embargoed Tue 9am ET: Company news in 8 words]").font = F_MUTED
    for r in range(5, plast + 1):
        pt.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        f = pt.cell(row=r, column=8, value=f'=IF(A{r}="","",A{r}+4)')
        f.number_format = "yyyy-mm-dd"
        for col in range(1, 13):
            pt.cell(row=r, column=col).border = BORDER
            pt.cell(row=r, column=col).alignment = WRAP
    pt.conditional_formatting.add(f"J5:J{plast}", CellIsRule(operator="equal", formula=['"Covered"'], fill=FILL_LIME))
    pt.auto_filter.ref = f"A4:L{plast}"

    # Summary block on the right of the tracker
    pt.column_dimensions["N"].width = 22
    pt.column_dimensions["O"].width = 10
    pt["N4"] = "Summary"
    pt["N4"].font = F_H
    pt["N4"].fill = FILL_H
    pt["O4"].fill = FILL_H
    summary = [("Pitches sent", f'=COUNTA(A5:A{plast})')] + [
        (o, f'=COUNTIF(J5:J{plast},"{o}")') for o in OUTCOMES.split(",")
    ] + [("Hit rate (covered / sent)", f'=IFERROR(COUNTIF(J5:J{plast},"Covered")/COUNTA(A5:A{plast}),0)')]
    for i, (label, formula) in enumerate(summary):
        a = pt.cell(row=5 + i, column=14, value=label)
        b = pt.cell(row=5 + i, column=15, value=formula)
        a.border = BORDER
        b.border = BORDER
        if "rate" in label:
            b.number_format = PCT
    pt.cell(row=5 + len(summary) + 1, column=14, value=FOOTER).font = F_MUTED

    for s in wb.worksheets:
        s.sheet_properties.tabColor = LIME if s.title != "Instructions" else INK
    wb.active = 0
    wb.save(path)
    return path


def build(out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    return [
        build_budget(out_dir / "marketing-budget-template.xlsx"),
        build_media_list(out_dir / "media-list-template.xlsx"),
    ]
